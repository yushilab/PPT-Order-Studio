"""Structural PPTX checks and optional local PDF rendering.

The inspection report never certifies actual text fit or visual quality.
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import tempfile
from pathlib import Path


def inspect_pptx(path: Path, min_font_size: int = 12) -> dict:
    from pptx import Presentation
    if min_font_size < 1:
        raise ValueError("min_font_size must be positive")
    issues = []
    try:
        prs = Presentation(str(path))
    except Exception as exc:
        return {"passed": False, "issues": [{"severity": "error", "code": "unreadable_pptx", "message": str(exc)}], "slide_count": 0, "visual_review_required": True}
    if not prs.slides:
        issues.append({"severity": "error", "code": "no_slides", "message": "Presentation has no slides"})
    tolerance = 12700  # one point; account for integer rounding on export
    for number, slide in enumerate(prs.slides, 1):
        meaningful = False

        def inspect_text_frame(frame, name):
            nonlocal meaningful
            value = frame.text.strip()
            if value:
                meaningful = True
            if any(token in value.lower() for token in ("lorem ipsum", "待补充", "占位文字", "[todo]", "[tbd]")):
                issues.append({"severity": "warning", "code": "placeholder", "slide": number, "shape": name, "message": "Potential placeholder text needs review"})
            for paragraph in frame.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip() and run.font.size is not None and run.font.size.pt < min_font_size-.1:
                        issues.append({"severity": "warning", "code": "small_font", "slide": number, "shape": name, "message": f"Explicit font size {run.font.size.pt:g}pt is below {min_font_size}pt"})

        def inspect_shape(shape, nested=False):
            nonlocal meaningful
            # Group children have group-space coordinates; inspect the outer
            # group's bounds rather than pretending those are slide-space.
            if not nested and (shape.left < -tolerance or shape.top < -tolerance or shape.left+shape.width > prs.slide_width+tolerance or shape.top+shape.height > prs.slide_height+tolerance):
                issues.append({"severity": "error", "code": "out_of_bounds", "slide": number, "shape": shape.name, "message": "Object extends beyond the slide boundary"})
            if shape.has_chart or shape.shape_type == 13:
                meaningful = True
            if shape.has_text_frame:
                inspect_text_frame(shape.text_frame, shape.name)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        inspect_text_frame(cell.text_frame, shape.name)
            if shape.shape_type == 6:
                for child in shape.shapes:
                    inspect_shape(child, nested=True)

        for shape in slide.shapes:
            inspect_shape(shape)
        if not meaningful:
            issues.append({"severity": "error", "code": "empty_slide", "slide": number, "message": "No text, picture or chart detected"})
    return {"passed": not any(i["severity"] == "error" for i in issues), "issues": issues, "slide_count": len(prs.slides), "visual_review_required": True, "limitations": ["No font installation or substitution certification", "No proof of text fit, overlap, factual accuracy or visual quality", "Groups, tables, charts and inherited font sizes require visual review"]}


def render_pptx(path: Path, out_dir: Path, backend: str = "auto") -> dict:
    """Export a local PDF when Office/LibreOffice is available.

    Native applications execute in a bounded subprocess. Existing input files
    are opened read-only. PDF pages can then be inspected with an optional viewer.
    """
    if backend not in {"auto", "libreoffice", "powerpoint"}:
        raise ValueError("backend must be auto, libreoffice or powerpoint")
    path, out_dir = Path(path).resolve(), Path(out_dir).resolve()
    if not path.is_file():
        raise FileNotFoundError(path)
    out_dir.mkdir(parents=True, exist_ok=True)
    destination = out_dir / (path.stem + ".pdf")
    if destination.exists():
        raise FileExistsError(f"PDF already exists: {destination}")
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable and os.name == "nt":
        for base in (os.environ.get("ProgramFiles"), os.environ.get("ProgramFiles(x86)")):
            candidate = Path(base or "") / "LibreOffice" / "program" / "soffice.exe"
            if base and candidate.is_file():
                executable = str(candidate)
                break
    selected = backend
    if backend == "auto":
        selected = "libreoffice" if executable else "powerpoint" if os.name == "nt" else "unavailable"
    if selected == "libreoffice" and not executable:
        return {"status": "unavailable", "backend": selected, "message": "LibreOffice was not found; install it or export PDF manually in PowerPoint/WPS"}
    if selected == "powerpoint" and os.name != "nt":
        return {"status": "unavailable", "backend": selected, "message": "Automatic PowerPoint export is supported on Windows only; export PDF manually"}
    if selected == "unavailable":
        return {"status": "unavailable", "backend": selected, "message": "Install LibreOffice or export a PDF manually in PowerPoint/WPS"}
    # Use an isolated output directory so stale PDFs cannot be mistaken for success.
    with tempfile.TemporaryDirectory(prefix="pptstudio-render-", dir=out_dir, ignore_cleanup_errors=True) as temporary:
        tmp = Path(temporary)
        pdf = tmp / (path.stem + ".pdf")
        powershell_input = None
        environment = None
        if selected == "libreoffice":
            profile = tmp / "profile"
            command = [executable, "-env:UserInstallation="+profile.as_uri(), "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(path)]
        else:
            powershell = shutil.which("powershell") or shutil.which("pwsh")
            if not powershell:
                return {"status": "unavailable", "backend": selected, "message": "PowerShell not found; export PDF manually"}
            # Paths travel in JSON, never shell interpolation. The script is fixed.
            config = tmp / "render.json"
            config.write_text(json.dumps({"input": str(path), "output": str(pdf)}), encoding="utf-8")
            powershell_input = '''
$ErrorActionPreference = 'Stop'
$settings = Get-Content -LiteralPath $env:PPTSTUDIO_RENDER_CONFIG -Raw -Encoding UTF8 | ConvertFrom-Json
$office = $null
$deck = $null
try {
    $office = New-Object -ComObject PowerPoint.Application
    $deck = $office.Presentations.Open($settings.input, $true, $false, $false)
    $deck.SaveAs($settings.output, 32)
} finally {
    if ($null -ne $deck) { $deck.Close(); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($deck) }
    if ($null -ne $office) { [void][Runtime.InteropServices.Marshal]::ReleaseComObject($office) }
}

'''
            environment = os.environ.copy()
            environment["PPTSTUDIO_RENDER_CONFIG"] = str(config)
            # Read fixed commands from stdin. No execution-policy changes and
            # no interpolation of customer-controlled paths into shell code.
            command = [powershell, "-NoProfile", "-NonInteractive", "-Command", "-"]
        try:
            result = subprocess.run(command, input=powershell_input, env=environment, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120, check=False)
        except subprocess.TimeoutExpired:
            return {"status": "error", "backend": selected, "message": "Rendering exceeded 120 seconds; export PDF manually and check for an Office dialog"}
        except OSError as exc:
            return {"status": "unavailable", "backend": selected, "message": str(exc)}
        if result.returncode != 0 or not pdf.exists():
            message = (result.stderr or result.stdout or "No PDF output produced")[-1800:]
            return {"status": "unavailable" if selected == "powerpoint" else "error", "backend": selected, "message": message, "manual_fallback": "Open the PPTX in the target PowerPoint/WPS and export PDF"}
        if destination.exists():
            raise FileExistsError(f"PDF already exists: {destination}")
        shutil.copy2(pdf, destination)
    return {"status": "success", "backend": selected, "pdf": str(destination), "visual_review_required": True}
