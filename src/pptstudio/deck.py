"""Deterministic, editable 16:9 PowerPoint layouts; no model or network calls."""
from __future__ import annotations

import json
import math
import re
from pathlib import Path
from typing import Any

DEFAULT_STYLE = {
    "schema_version": 1, "name": "Forest editorial", "font": "Microsoft YaHei",
    "background": "F7F9FC", "foreground": "17342C", "primary": "1C6951",
    "accent": "C49A41", "muted": "61746D", "title_size": 32,
    "body_size": 20, "min_font_size": 12,
}
LAYOUTS = {"cover", "bullets", "metrics", "chart", "image", "timeline", "closing"}
WIDTH, HEIGHT = 13.333333, 7.5


def _json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def _inside(order: Path, value: str) -> Path:
    """Resolve an order-relative asset, rejecting traversal and symlink escapes."""
    p = Path(value)
    if not value or p.is_absolute() or re.match(r"^[A-Za-z]+:", value):
        raise ValueError("image must be an order-relative file path")
    if ".." in value.replace("\\", "/").split("/"):
        raise ValueError("image path must not contain '..'")
    root = order.resolve()
    candidate = (root / p).resolve()
    if not candidate.is_relative_to(root):
        raise ValueError("image path escapes the order directory")
    if not candidate.is_file():
        raise ValueError(f"image file not found: {value}")
    return candidate


def validate_order(order: Path, slide_ids: list[str] | None = None, outline: bool = False) -> list[dict]:
    """Check the portable order schema and conservative layout content limits.

    This validates structure and supported content, not truth or visual fit.
    """
    order = Path(order)
    issues: list[dict] = []

    def issue(code: str, message: str, sid: str | None = None, severity="error"):
        row = {"severity": severity, "code": code, "message": message}
        if sid:
            row["slide_id"] = sid
        issues.append(row)

    loaded = {}
    for name in ("brief", "deck", "style"):
        try:
            value = _json(order / f"{name}.json")
            if not isinstance(value, dict):
                raise ValueError("root must be an object")
            loaded[name] = value
        except (OSError, ValueError, TypeError) as exc:
            issue("invalid_json", f"{name}.json: {exc}")
    if len(loaded) != 3:
        return issues
    brief, deck, style = loaded["brief"], loaded["deck"], {**DEFAULT_STYLE, **loaded["style"]}
    for name, data in loaded.items():
        if data.get("schema_version") != 1:
            issue("schema_version", f"{name}.json requires schema_version: 1")
    count = brief.get("page_count")
    if isinstance(count, bool) or not isinstance(count, int) or not 1 <= count <= 200:
        issue("page_count", "brief.page_count must be an integer between 1 and 200")
    if not isinstance(brief.get("title"), str) or not brief["title"].strip():
        issue("brief_title", "brief.title must be a non-empty string")
    for key in ("background", "foreground", "primary", "accent", "muted"):
        if not isinstance(style[key], str) or not re.fullmatch(r"[0-9a-fA-F]{6}", style[key]):
            issue("style_color", f"style.{key} requires six hex digits, without '#'")
    if not isinstance(style["font"], str) or not style["font"].strip():
        issue("style_font", "style.font must name a font installed on the target machine")
    for key in ("title_size", "body_size", "min_font_size"):
        value = style[key]
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not 10 <= value <= 60:
            issue("style_size", f"style.{key} must be a finite number between 10 and 60")
    if isinstance(style["body_size"], (int, float)) and isinstance(style["min_font_size"], (int, float)):
        if style["body_size"] < style["min_font_size"]:
            issue("style_size", "body_size must not be below min_font_size")
    source_ids: set[str] = set()
    source_status = {}
    facts: dict[str, dict] = {}
    manifest = order / "sources.json"
    if manifest.exists():
        try:
            files = _json(manifest).get("files", [])
            if not isinstance(files, list):
                raise ValueError("files must be a list")
            for record in files:
                if isinstance(record, dict) and isinstance(record.get("id"), str):
                    source_ids.add(record["id"])
                    source_status[record["id"]] = record.get("status")
        except (OSError, ValueError, AttributeError) as exc:
            issue("sources_schema", f"sources.json: {exc}")
    if (order / "facts.json").exists():
        try:
            rows = _json(order / "facts.json").get("facts", [])
            if not isinstance(rows, list):
                raise ValueError("facts must be a list")
            for fact in rows:
                if not isinstance(fact, dict) or not isinstance(fact.get("id"), str) or not fact["id"].strip():
                    issue("facts_schema", "Each fact needs a non-empty string id")
                    continue
                if fact["id"] in facts or fact["id"] in source_ids:
                    issue("facts_schema", f"Duplicate or ambiguous fact ID: {fact['id']}")
                facts[fact["id"]] = fact
        except (OSError, ValueError, AttributeError) as exc:
            issue("facts_schema", f"facts.json: {exc}")

    def check_ref(ref: str, sid: str, chain: frozenset = frozenset()):
        base = ref.split("#", 1)[0]
        if base in source_ids:
            if source_status.get(base) != "ok":
                issue("source_review", f"Source {base} requires review: {source_status.get(base)}", sid, "warning")
        elif base in facts:
            if base in chain:
                issue("fact_cycle", f"Circular fact reference: {base}", sid)
                return
            refs = facts[base].get("source_refs", [])
            if not isinstance(refs, list) or any(not isinstance(r, str) for r in refs):
                issue("facts_schema", f"Fact {base}.source_refs must be a string list", sid)
            elif not refs:
                issue("missing_fact_sources", f"Fact {base} has no source_refs; verify manually", sid, "warning")
            else:
                for nested in refs:
                    check_ref(nested, sid, chain | {base})
        else:
            issue("unknown_source", f"Unknown source or fact ID: {base}", sid)
    slides = deck.get("slides")
    if not isinstance(slides, list) or not slides:
        issue("slides_empty", "deck.slides must contain at least one slide")
        return issues
    if isinstance(count, int) and len(slides) != count:
        issue("page_count_mismatch", f"Expected {count} slides, found {len(slides)}")
    if slide_ids is not None:
        if not slide_ids or len(slide_ids) != len(set(slide_ids)):
            issue("selection", "Select unique, non-empty slide IDs")
        available = {s.get("id") for s in slides if isinstance(s, dict) and isinstance(s.get("id"), str)}
        for selected in slide_ids:
            if selected not in available:
                issue("selection", f"Unknown slide ID: {selected}")
    seen = set()
    for index, slide in enumerate(slides, 1):
        if not isinstance(slide, dict):
            issue("slide_schema", f"Slide {index} must be an object")
            continue
        sid = slide.get("id")
        if not isinstance(sid, str) or not re.fullmatch(r"S\d{2,4}", sid):
            issue("slide_id", f"Slide {index} needs an ID like S01")
            sid = f"slide-{index}"
        if sid in seen:
            issue("duplicate_id", f"Duplicate slide ID: {sid}", sid)
        seen.add(sid)
        layout = slide.get("layout")
        if not isinstance(layout, str) or layout not in LAYOUTS:
            issue("layout", f"Unsupported layout: {layout}", sid)
            layout = ""
        for field, limit in (("title", 100), ("message", 240)):
            value = slide.get(field, "")
            if not isinstance(value, str) or (field == "title" and not value.strip()):
                issue("slide_text", f"{field} must be a {'non-empty ' if field == 'title' else ''}string", sid)
            elif len(value) > limit:
                issue("content_density", f"{field} exceeds {limit} characters; shorten or split the page", sid)
        refs = slide.get("source_refs", [])
        if not isinstance(refs, list) or any(not isinstance(r, str) for r in refs):
            issue("source_refs", "source_refs must be a list of source IDs", sid)
        else:
            for ref in refs:
                check_ref(ref, sid)
            if not refs and layout not in {"cover", "closing"}:
                issue("missing_sources", "No source_refs; verify this page's claims manually", sid, "warning")
        if outline or (slide_ids is not None and sid not in slide_ids):
            continue
        body = slide.get("body", [])
        if not isinstance(body, list) or any(not isinstance(t, str) for t in body):
            issue("body", "body must be a list of strings", sid)
        elif len(body) > 6 or any(len(t) > 180 for t in body) or sum(map(len, body)) > 650:
            issue("content_density", "body allows up to six items, 180 characters per item and 650 total", sid)
        if layout == "bullets" and not body:
            issue("body_empty", "bullets layout needs body items", sid)
        if layout == "metrics":
            metrics = slide.get("metrics")
            if not isinstance(metrics, list) or not 1 <= len(metrics) <= 4:
                issue("metrics", "metrics layout needs one to four metric objects", sid)
            else:
                for metric in metrics:
                    if not isinstance(metric, dict) or any(not isinstance(metric.get(k), str) or not metric[k].strip() for k in ("label", "value")):
                        issue("metrics", "Each metric needs non-empty string label and value", sid)
                    elif len(metric["label"]) > 60 or len(metric["value"]) > 24:
                        issue("content_density", "Metric label/value exceeds 60/24 characters", sid)
        if layout == "chart":
            chart = slide.get("chart")
            if not isinstance(chart, dict):
                issue("chart", "chart layout needs a chart object", sid)
            else:
                categories, series = chart.get("categories"), chart.get("series")
                if not isinstance(categories, list) or not 1 <= len(categories) <= 12 or any(not isinstance(c, str) or len(c) > 24 for c in categories):
                    issue("chart_categories", "chart.categories needs one to twelve labels, up to 24 characters each", sid)
                if not isinstance(series, list) or not 1 <= len(series) <= 3:
                    issue("chart_series", "chart.series needs one to three series", sid)
                else:
                    for row in series:
                        if not isinstance(row, dict) or not isinstance(row.get("name"), str):
                            issue("chart_series", "Each series needs name and values", sid)
                            continue
                        vals = row.get("values")
                        if not isinstance(vals, list) or not isinstance(categories, list) or len(vals) != len(categories) or any(isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) for v in vals):
                            issue("chart_values", "Each series needs finite numeric values matching the category count", sid)
        if layout == "image":
            image = slide.get("image")
            if not isinstance(image, str):
                issue("image", "image layout requires an order-relative image path", sid)
            else:
                try:
                    asset = _inside(order, image)
                    from PIL import Image
                    with Image.open(asset) as img:
                        img.verify()
                except (ValueError, OSError, ImportError) as exc:
                    issue("image", str(exc), sid)
            if not isinstance(slide.get("caption", ""), str) or len(slide.get("caption", "")) > 350:
                issue("caption", "caption must be a string up to 350 characters", sid)
        if layout == "timeline":
            events = slide.get("events")
            if not isinstance(events, list) or not 2 <= len(events) <= 6:
                issue("timeline", "timeline needs two to six event objects", sid)
            else:
                for event in events:
                    if not isinstance(event, dict) or any(not isinstance(event.get(k), str) or not event[k].strip() for k in ("year", "text")):
                        issue("timeline", "Each event needs string year and text", sid)
                    elif len(event["year"]) > 16 or len(event["text"]) > 140:
                        issue("content_density", "Event year/text exceeds 16/140 characters", sid)
    return issues


def build_deck(order: Path, output: Path, slide_ids: list[str] | None = None, draft: bool = False) -> dict:
    """Build editable slides. Gate approvals are enforced by the calling CLI."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    order, output = Path(order).resolve(), Path(output).resolve()
    errors = [i for i in validate_order(order, slide_ids=slide_ids) if i["severity"] == "error"]
    if errors:
        raise ValueError("Order validation failed: " + "; ".join(i["message"] for i in errors[:8]))
    if not output.is_relative_to(order):
        raise ValueError("PPTX output must stay inside the order directory")
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite an existing presentation: {output}")
    slides = _json(order / "deck.json")["slides"]
    if slide_ids is not None:
        if not slide_ids or len(slide_ids) != len(set(slide_ids)):
            raise ValueError("Select unique, non-empty slide IDs")
        by_id = {s["id"]: s for s in slides}
        unknown = [sid for sid in slide_ids if sid not in by_id]
        if unknown:
            raise ValueError("Unknown slide IDs: " + ", ".join(unknown))
        slides = [by_id[sid] for sid in slide_ids]
    brief, style = _json(order / "brief.json"), {**DEFAULT_STYLE, **_json(order / "style.json")}
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(WIDTH), Inches(HEIGHT)
    prs.core_properties.title = brief["title"]
    prs.core_properties.subject = "DRAFT — customer confirmation required" if draft else "PPT Order Studio"
    prs.core_properties.author = "PPT Order Studio"

    def rgb(key):
        return RGBColor.from_string(style.get(key, key))

    def text(slide, value, x, y, w, h, size, color="foreground", bold=False, align=PP_ALIGN.LEFT):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.02)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        tf.vertical_anchor = MSO_ANCHOR.TOP
        # A conservative approximation prevents extreme strings from silently
        # shrinking below the user's minimum. Actual font metrics need review.
        weighted = sum(1.0 if ord(ch) > 255 else 0.53 for ch in str(value))
        chosen = float(size)
        while chosen > style["min_font_size"]:
            chars_per_line = max(1, w * 72 / chosen)
            lines = max(1, math.ceil(weighted / chars_per_line))
            if lines * chosen * 1.25 / 72 <= h:
                break
            chosen -= 1
        for n, line in enumerate(str(value).split("\n")):
            p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_before = Pt(0)
            p.space_after = Pt(4)
            p.line_spacing = 1.12
            run = p.add_run()
            run.text = line
            run.font.name = style["font"]
            run.font.size = Pt(max(style["min_font_size"], chosen))
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
        return shape

    def rect(slide, x, y, w, h, color, radius=False):
        kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
        shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(color)
        shape.line.fill.background()
        return shape

    def base(item, number):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb("background")
        text(slide, brief["title"][:90], .62, 7.05, 10.8, .28, 12, "muted")
        text(slide, f"{item['id']}   /   {number:02d}", 11.75, 7.05, .95, .28, 12, "muted", align=PP_ALIGN.RIGHT)
        if draft:
            text(slide, "DRAFT · 待确认", 10.5, .22, 2.15, .3, 12, "muted", align=PP_ALIGN.RIGHT)
        if item["layout"] not in {"cover", "closing"}:
            text(slide, item["title"], .62, .58, 12.03, .75, style["title_size"], bold=True)
            if item.get("message"):
                text(slide, item["message"], .66, 1.4, 11.95, .65, style["body_size"], "muted")
        note = f"Slide ID: {item['id']}\nLayout: {item['layout']}\nSources: " + ", ".join(item.get("source_refs", []))
        if draft:
            note += "\nDRAFT: not approved for delivery."
        slide.notes_slide.notes_text_frame.text = note
        return slide

    for index, item in enumerate(slides, 1):
        slide = base(item, index)
        layout = item["layout"]
        if layout == "cover":
            text(slide, style["name"].upper(), .68, 1.04, 8.5, .45, 14, "primary", bold=True)
            text(slide, item["title"], .62, 2.0, 9.6, 2.6, 47, bold=True)
            text(slide, item.get("message", ""), .68, 4.95, 10.8, 1.0, 24, "muted")
            text(slide, "01", 10.5, 1.7, 2.1, 2.0, 60, "accent", bold=True, align=PP_ALIGN.RIGHT)
        elif layout == "bullets":
            body = item["body"]
            row_h = 4.45 / len(body)
            for n, line in enumerate(body):
                y = 2.2 + n * row_h
                text(slide, f"{n+1:02d}", .68, y, .7, row_h - .1, 23, "primary", bold=True)
                text(slide, line, 1.55, y, 10.9, row_h - .12, style["body_size"])
        elif layout == "metrics":
            metrics = item["metrics"]
            gap = .25
            w = (12.0 - gap * (len(metrics)-1)) / len(metrics)
            for n, metric in enumerate(metrics):
                x = .65 + n * (w + gap)
                rect(slide, x, 2.65, w, 2.75, "primary", radius=True)
                text(slide, metric["value"], x+.24, 3.1, w-.48, 1.0, 43, "background", bold=True)
                text(slide, metric["label"], x+.24, 4.28, w-.48, .9, 20, "background")
            if item.get("body"):
                text(slide, "\n".join(item["body"]), .7, 5.8, 11.9, .95, 18, "muted")
        elif layout == "chart":
            data = CategoryChartData()
            data.categories = item["chart"]["categories"]
            for series in item["chart"]["series"]:
                data.add_series(series["name"], series["values"])
            has_sidebar = bool(item.get("body"))
            cw = 8.75 if has_sidebar else 11.9
            chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.65), Inches(2.25), Inches(cw), Inches(4.2), data).chart
            chart.has_legend = len(item["chart"]["series"]) > 1
            if chart.has_legend:
                chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                chart.legend.font.name, chart.legend.font.size = style["font"], Pt(12)
            chart.chart_style = 10
            for axis in (chart.category_axis, chart.value_axis):
                axis.tick_labels.font.name = style["font"]
                axis.tick_labels.font.size = Pt(12)
                axis.tick_labels.font.color.rgb = rgb("muted")
            chart.value_axis.has_major_gridlines = True
            for n, series in enumerate(chart.series):
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = rgb(("primary", "accent", "muted")[n])
                series.format.line.fill.background()
            chart.plots[0].gap_width = 95
            chart.plots[0].has_data_labels = True
            labels = chart.plots[0].data_labels
            labels.font.name, labels.font.size = style["font"], Pt(12)
            labels.font.color.rgb = rgb("foreground")
            labels.position = XL_LABEL_POSITION.OUTSIDE_END
            if has_sidebar:
                text(slide, "\n\n".join(item["body"]), 9.8, 2.55, 2.8, 3.7, 19)
        elif layout == "image":
            from PIL import Image
            asset = _inside(order, item["image"])
            with Image.open(asset) as image:
                iw, ih = image.size
            box_w = 8.1 if item.get("caption") else 12.0
            box_h = 4.45
            scale = min(box_w / iw, box_h / ih)
            pw, ph = iw*scale, ih*scale
            slide.shapes.add_picture(str(asset), Inches(.65+(box_w-pw)/2), Inches(2.2+(box_h-ph)/2), width=Inches(pw), height=Inches(ph))
            if item.get("caption"):
                text(slide, item["caption"], 9.18, 2.7, 3.42, 3.8, 20)
        elif layout == "timeline":
            events = item["events"]
            row_h = 4.35/len(events)
            for n, event in enumerate(events):
                y = 2.2+n*row_h
                text(slide, event["year"], .7, y, 1.65, row_h-.1, 21, "primary", bold=True)
                rect(slide, 2.6, y+.1, .06, max(.12, row_h-.22), "accent")
                text(slide, event["text"], 3.05, y, 9.4, row_h-.1, style["body_size"])
        elif layout == "closing":
            text(slide, item["title"], .68, 2.05, 11.7, 1.7, 45, bold=True)
            text(slide, item.get("message", ""), .72, 4.0, 11.6, 1.1, 25, "primary")
            if item.get("body"):
                text(slide, "\n".join(item["body"]), .72, 5.55, 11.5, 1.0, 18, "muted")

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(output.name + ".tmp")
    try:
        prs.save(str(temporary))
        temporary.replace(output)
    finally:
        if temporary.exists():
            temporary.unlink()
    return {"path": str(output), "slide_ids": [s["id"] for s in slides], "slide_count": len(slides), "draft": draft}
