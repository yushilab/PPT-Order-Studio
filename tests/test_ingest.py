from __future__ import annotations

import importlib.util
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptstudio.ingest import ingest_order, source_excerpt


class IngestTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.order = Path(self.temp.name) / "中文 order"
        (self.order / "sources").mkdir(parents=True)

    def tearDown(self):
        self.temp.cleanup()

    def write(self, name, text):
        path = self.order / "sources" / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")
        return path

    def test_incremental_cache_and_stable_id_and_removed(self):
        path = self.write("材料.md", "# 客流\n2025 年：125 万人次")
        first = ingest_order(self.order)
        record = first["files"][0]
        self.assertEqual(record["status"], "ok")
        self.assertEqual(record["blocks"][0]["locator"], "lines 1-2")
        self.assertIn("125 万人次", record["blocks"][0]["text"])
        text_path = self.order / record["text_path"]
        timestamp = text_path.stat().st_mtime_ns
        again = ingest_order(self.order)
        self.assertEqual(again["unchanged"], [record["id"]])
        self.assertEqual(text_path.stat().st_mtime_ns, timestamp)
        path.write_text("# 客流\n2025 年：130 万人次", encoding="utf-8")
        changed = ingest_order(self.order)
        self.assertEqual(changed["files"][0]["id"], record["id"])
        self.assertNotEqual(changed["files"][0]["sha256"], record["sha256"])
        self.assertEqual(changed["changed"], [record["id"]])
        path.unlink()
        removed = ingest_order(self.order)
        self.assertEqual(removed["removed"], [record["id"]])
        self.assertEqual(removed["files"], [])

    def test_corrupt_extracted_cache_is_rebuilt(self):
        self.write("facts.txt", "actual content")
        first = ingest_order(self.order)["files"][0]
        (self.order / first["text_path"]).write_text("malicious replacement", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "modified"):
            source_excerpt(self.order, first["id"])
        rebuilt = ingest_order(self.order)
        self.assertEqual(rebuilt["changed"], [first["id"]])
        self.assertIn("actual content", source_excerpt(self.order, first["id"]))

    def test_csv_preserves_values_and_cell_references(self):
        self.write("data.csv", '项目,变化,备注\n景区,-0.25,"带,逗号"\n')
        blocks = ingest_order(self.order)["files"][0]["blocks"]
        self.assertEqual(blocks[1]["locator"], "row 2")
        self.assertIn('B2="-0.25"', blocks[1]["text"])
        self.assertIn('C2="带,逗号"', blocks[1]["text"])

    def test_utf16_and_legacy_formats(self):
        (self.order / "sources" / "utf16.txt").write_bytes("收入单位：万元".encode("utf-16"))
        (self.order / "sources" / "legacy.xls").write_bytes(b"not-xlsx")
        records = {Path(item["path"]).name: item for item in ingest_order(self.order)["files"]}
        self.assertEqual(records["legacy.xls"]["status"], "unsupported")
        self.assertIn(".xlsx", records["legacy.xls"]["warnings"][0])
        self.assertIn("万元", records["utf16.txt"]["blocks"][0]["text"])

    def test_excerpt_bounds_and_unknown_id(self):
        self.write("long.md", "hello\n" * 100)
        record = ingest_order(self.order)["files"][0]
        excerpt = source_excerpt(self.order, record["id"], limit=25)
        self.assertIn("character range 0:25", excerpt)
        self.assertIn("More available: start=25", excerpt)
        for start, limit in ((-1, 1), (0, 0), (0, 2_000_001)):
            with self.assertRaises(ValueError):
                source_excerpt(self.order, record["id"], start, limit)
        with self.assertRaisesRegex(ValueError, "Unknown source"):
            source_excerpt(self.order, "does-not-exist")

    def test_excerpt_refuses_stale_source_until_ingest(self):
        source = self.write("facts.md", "old fact")
        record = ingest_order(self.order)["files"][0]
        source.write_text("new fact", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "Source changed"):
            source_excerpt(self.order, record["id"])
        ingest_order(self.order)
        self.assertIn("new fact", source_excerpt(self.order, record["id"]))

    def test_missing_optional_reader_errors_are_retried(self):
        self.write("reader.xlsx", "placeholder")
        missing = ImportError("reader absent", name="openpyxl")
        with patch("pptstudio.ingest._extract", side_effect=missing):
            first = ingest_order(self.order)
        self.assertEqual(first["files"][0]["status"], "error")
        self.assertIn("[extract]", first["files"][0]["warnings"][0])
        with patch("pptstudio.ingest._extract", return_value=("ok", [{"locator": "A1", "text": "now readable"}], [], [])) as reader:
            second = ingest_order(self.order)
        reader.assert_called_once()
        self.assertEqual(second["files"][0]["status"], "ok")
        self.assertEqual(second["changed"], [first["files"][0]["id"]])

    def test_manifest_path_escape_is_not_read(self):
        self.write("safe.txt", "safe")
        record = ingest_order(self.order)["files"][0]
        manifest = json.loads((self.order / "sources.json").read_text(encoding="utf-8"))
        manifest["files"][0]["text_path"] = "../outside.txt"
        (self.order / "sources.json").write_text(json.dumps(manifest), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "escapes"):
            source_excerpt(self.order, record["id"])
        rebuilt = ingest_order(self.order)
        self.assertEqual(rebuilt["changed"], [record["id"]])

    def test_source_symlink_is_rejected(self):
        target = self.order / "outside.txt"
        target.write_text("outside", encoding="utf-8")
        link = self.order / "sources" / "linked.txt"
        try:
            link.symlink_to(target)
        except (OSError, NotImplementedError):
            self.skipTest("Symlinks unavailable for this user")
        with self.assertRaisesRegex(ValueError, "link"):
            ingest_order(self.order)

    @unittest.skipUnless(importlib.util.find_spec("docx"), "python-docx optional")
    def test_docx_paragraph_table_and_image(self):
        from docx import Document
        document = Document()
        document.add_paragraph("汇报主题")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "125 万人次"
        document.add_paragraph("表格之后的结论")
        document.save(self.order / "sources" / "source.docx")
        record = ingest_order(self.order)["files"][0]
        self.assertEqual(record["status"], "ok")
        self.assertTrue(any(item["locator"] == "paragraph 1" and "主题" in item["text"] for item in record["blocks"]))
        self.assertTrue(any(item["locator"] == "table 1 row 1" and "125" in item["text"] for item in record["blocks"]))
        self.assertEqual([item["locator"] for item in record["blocks"][:3]], ["paragraph 1", "table 1 row 1", "paragraph 2"])

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl optional")
    def test_excel_formula_missing_cache_is_explicit(self):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.title = "经营汇总"
        sheet["A1"] = "单位：万元"
        sheet["B2"] = "=1+2"
        sheet.row_dimensions[2].hidden = True
        sheet.merge_cells("A3:B3")
        book.save(self.order / "sources" / "source.xlsx")
        record = ingest_order(self.order)["files"][0]
        self.assertEqual(record["status"], "ok")
        self.assertTrue(any(item["locator"] == "sheet 经营汇总!B2" and "cached_value=None" in item["text"] for item in record["blocks"]))
        self.assertTrue(any("no cached result" in text for text in record["warnings"]))
        self.assertTrue(any("hidden content" in text for text in record["warnings"]))
        self.assertTrue(any(item["locator"].endswith("merged ranges") for item in record["blocks"]))

    @unittest.skipUnless(importlib.util.find_spec("pypdf"), "pypdf optional")
    def test_blank_pdf_is_flagged_without_fabricated_text(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=720, height=480)
        with (self.order / "sources" / "scan.pdf").open("wb") as handle:
            writer.write(handle)
        record = ingest_order(self.order)["files"][0]
        self.assertEqual(record["status"], "needs_ocr")
        self.assertEqual(record["blocks"][0], {"locator": "page 1", "text": ""})

    @unittest.skipUnless(importlib.util.find_spec("PIL"), "Pillow optional")
    def test_image_metadata_is_not_ocr_and_asset_cache_is_checked(self):
        from PIL import Image
        Image.new("RGB", (20, 10), "white").save(self.order / "sources" / "image.png")
        record = ingest_order(self.order)["files"][0]
        self.assertEqual(record["status"], "needs_ocr")
        self.assertIn("20x10", record["blocks"][0]["text"])
        self.assertIn("No OCR", record["blocks"][0]["text"])
        self.assertEqual(record["assets"][0]["path"], "sources/image.png")


if __name__ == "__main__":
    unittest.main()
