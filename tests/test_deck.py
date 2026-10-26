"""Checks exercise editable objects, content limits and filesystem invariants."""
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches, Pt

from pptstudio.deck import DEFAULT_STYLE, build_deck, validate_order
from pptstudio.quality import inspect_pptx, render_pptx


class DeckTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.order = Path(self.temporary.name) / "order"
        self.order.mkdir()
        self.slides = [
            {"id": "S01", "layout": "cover", "title": "Annual review", "message": "Results and next steps", "source_refs": []},
            {"id": "S02", "layout": "bullets", "title": "Three priorities", "message": "Make the workflow repeatable", "body": ["Validate sources", "Confirm representative samples", "Revise only affected pages"], "source_refs": ["src-demo"]},
            {"id": "S03", "layout": "metrics", "title": "Measured outcomes", "metrics": [{"label": "Orders", "value": "24"}, {"label": "Delivery rate", "value": "98%"}], "source_refs": ["src-demo"]},
            {"id": "S04", "layout": "chart", "title": "Quarterly deliveries", "chart": {"categories": ["Q1", "Q2", "Q3"], "series": [{"name": "Orders", "values": [4, 8, 12]}]}, "body": ["Fictional demonstration data"], "source_refs": ["src-demo"]},
            {"id": "S05", "layout": "timeline", "title": "Delivery process", "events": [{"year": "Plan", "text": "Confirm the page outline"}, {"year": "Sample", "text": "Confirm design"}, {"year": "Deliver", "text": "Inspect the final PPTX"}], "source_refs": ["src-demo"]},
            {"id": "S06", "layout": "closing", "title": "Ready for review", "message": "Human approval is required", "source_refs": []},
        ]
        self.write("brief.json", {"schema_version": 1, "title": "Demo order", "page_count": len(self.slides)})
        self.write("style.json", DEFAULT_STYLE)
        self.write("sources.json", {"schema_version": 1, "files": [{"id": "src-demo", "status": "ok"}]})
        self.save_slides()

    def tearDown(self):
        self.temporary.cleanup()

    def write(self, name, value):
        (self.order / name).write_text(json.dumps(value), encoding="utf-8")

    def save_slides(self):
        self.write("deck.json", {"schema_version": 1, "slides": self.slides})

    def test_native_editable_text_chart_and_notes(self):
        target = self.order / "build" / "demo.pptx"
        result = build_deck(self.order, target, draft=True)
        self.assertEqual(result["slide_count"], 6)
        prs = Presentation(target)
        self.assertEqual(len(prs.slides), 6)
        self.assertTrue(any(s.has_chart for s in prs.slides[3].shapes))
        self.assertTrue(any(s.has_text_frame and "Three priorities" in s.text for s in prs.slides[1].shapes))
        self.assertIn("src-demo", prs.slides[3].notes_slide.notes_text_frame.text)
        self.assertIn("DRAFT", prs.slides[3].notes_slide.notes_text_frame.text)
        report = inspect_pptx(target)
        self.assertTrue(report["passed"], report)
        self.assertTrue(report["visual_review_required"])

    def test_subset_can_build_before_unselected_body_is_written(self):
        self.slides[1]["body"] = []
        self.save_slides()
        self.assertTrue(any(i["code"] == "body_empty" for i in validate_order(self.order)))
        result = build_deck(self.order, self.order / "sample.pptx", ["S01", "S03"])
        self.assertEqual(result["slide_ids"], ["S01", "S03"])

    def test_selected_incomplete_page_rejected(self):
        self.slides[1]["body"] = []
        self.save_slides()
        with self.assertRaises(ValueError):
            build_deck(self.order, self.order / "sample.pptx", ["S02"])

    def test_outline_allows_unfinished_layout_content(self):
        self.slides[1]["body"] = []
        self.slides[3].pop("chart")
        self.save_slides()
        self.assertFalse(any(i["severity"] == "error" for i in validate_order(self.order, outline=True)))
        self.assertTrue(any(i["severity"] == "error" for i in validate_order(self.order)))

    def test_locator_and_fact_references_are_resolved(self):
        self.write("facts.json", {"facts": [{"id": "F01", "text": "Quarterly deliveries", "source_refs": ["src-demo#table:1"]}]})
        self.slides[1]["source_refs"] = ["F01", "src-demo#line:4"]
        self.save_slides()
        self.assertFalse(any(i["severity"] == "error" for i in validate_order(self.order)))

    def test_fact_reference_cycle_is_reported(self):
        self.write("facts.json", {"facts": [{"id": "F01", "source_refs": ["F02"]}, {"id": "F02", "source_refs": ["F01"]}]})
        self.slides[1]["source_refs"] = ["F01"]
        self.save_slides()
        self.assertTrue(any(i["code"] == "fact_cycle" for i in validate_order(self.order)))

    def test_overwrite_refused(self):
        output = self.order / "result.pptx"
        build_deck(self.order, output)
        original = output.read_bytes()
        with self.assertRaises(FileExistsError):
            build_deck(self.order, output)
        self.assertEqual(original, output.read_bytes())

    def test_output_escape_refused(self):
        with self.assertRaises(ValueError):
            build_deck(self.order, self.order.parent / "escaped.pptx")

    def test_invalid_chart_values_rejected_without_partial_file(self):
        self.slides[3]["chart"]["series"][0]["values"] = [1, float("nan"), 3]
        self.save_slides()
        output = self.order / "invalid.pptx"
        with self.assertRaises(ValueError):
            build_deck(self.order, output)
        self.assertFalse(output.exists())

    def test_unknown_source_rejected(self):
        self.slides[1]["source_refs"] = ["unknown"]
        self.save_slides()
        self.assertTrue(any(i["code"] == "unknown_source" for i in validate_order(self.order)))

    def test_duplicate_slide_ids_rejected(self):
        self.slides[1]["id"] = "S01"
        self.save_slides()
        self.assertTrue(any(i["code"] == "duplicate_id" for i in validate_order(self.order)))

    def test_malformed_layout_is_reported_not_crashed(self):
        self.slides[1]["layout"] = {"bad": "object"}
        self.save_slides()
        self.assertTrue(any(i["code"] == "layout" for i in validate_order(self.order)))

    def test_content_density_is_reported(self):
        self.slides[1]["body"] = ["A"*181]
        self.save_slides()
        self.assertTrue(any(i["code"] == "content_density" for i in validate_order(self.order)))

    def test_exact_page_budget_enforced(self):
        self.write("brief.json", {"schema_version": 1, "title": "Demo", "page_count": 20})
        self.assertTrue(any(i["code"] == "page_count_mismatch" for i in validate_order(self.order)))

    def test_image_layout_embeds_local_asset(self):
        from PIL import Image
        path = self.order / "assets" / "figure.png"
        path.parent.mkdir()
        Image.new("RGB", (240, 120), "green").save(path)
        self.slides[1] = {"id": "S02", "layout": "image", "title": "Reference image", "image": "assets/figure.png", "caption": "A locally supplied figure", "source_refs": ["src-demo"]}
        self.save_slides()
        target = self.order / "image.pptx"
        build_deck(self.order, target, ["S02"])
        self.assertTrue(any(s.shape_type == 13 for s in Presentation(target).slides[0].shapes))

    def test_image_path_escape_rejected(self):
        self.slides[1] = {"id": "S02", "layout": "image", "title": "Escape", "image": "../secret.png", "source_refs": []}
        self.save_slides()
        self.assertTrue(any(i["code"] == "image" for i in validate_order(self.order)))

    def test_bad_pptx_reports_failure(self):
        bad = self.order / "bad.pptx"
        bad.write_text("not a zip")
        self.assertFalse(inspect_pptx(bad)["passed"])

    def test_out_of_bounds_shape_fails_structural_qa(self):
        target = self.order / "bounds.pptx"
        build_deck(self.order, target, ["S01"])
        prs = Presentation(target)
        shape = prs.slides[0].shapes.add_textbox(Inches(14), Inches(1), Inches(1), Inches(1))
        shape.text = "Outside"
        shape.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
        prs.save(target)
        report = inspect_pptx(target)
        self.assertFalse(report["passed"])
        self.assertTrue(any(i["code"] == "out_of_bounds" for i in report["issues"]))

    def test_missing_renderer_returns_unavailable(self):
        target = self.order / "render.pptx"
        build_deck(self.order, target, ["S01"])
        with patch("pptstudio.quality.shutil.which", return_value=None):
            result = render_pptx(target, self.order / "previews", backend="powerpoint")
        self.assertEqual(result["status"], "unavailable")

    def test_grouped_manual_content_is_not_reported_empty(self):
        target = self.order / "grouped.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        item = group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        item.text = "Manually grouped content"
        prs.save(target)
        self.assertTrue(inspect_pptx(target)["passed"])

    def test_renderer_refuses_existing_pdf_before_launching_office(self):
        target = self.order / "existing.pptx"
        build_deck(self.order, target, ["S01"])
        out_dir = self.order / "pdf"
        out_dir.mkdir()
        (out_dir / "existing.pdf").write_bytes(b"original")
        with patch("pptstudio.quality.subprocess.run") as runner:
            with self.assertRaises(FileExistsError):
                render_pptx(target, out_dir)
            runner.assert_not_called()


if __name__ == "__main__":
    unittest.main()
