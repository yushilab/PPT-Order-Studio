"""Bridge integrity tests using explicitly synthetic upstream-format fixtures.

These tests create a real PPTX ZIP and tiny test images, but do not call an AI,
browser, renderer, LibreOffice, or certify any visual quality. Synthetic review
records exist only inside these isolated test orders.
"""
import copy
import importlib.util
import json
import shutil
import subprocess
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

from pptstudio import core
from pptstudio import pptagent_bridge as bridge
from pptstudio.design import active_design, switch_design


@pytest.fixture
def order(tmp_path):
    root = tmp_path / "synthetic-order"
    core.init_order(root, "Synthetic bridge test, not a customer delivery", 3, engine="pptagent")
    core.write_json(root / "deck.json", {"schema_version": 1, "slides": [
        {"id": f"S{number:02d}", "title": f"Test {number}", "message": "Synthetic fixture",
         "layout": "custom", "body": ["Test data"], "source_refs": []}
        for number in range(1, 4)
    ]})
    return root


def fixture_result(order, selected=None, *, stage="sample", draft=False):
    if selected is None:
        selected = ([slide["id"] for slide in core.read_json(order / "deck.json")["slides"]]
                    if stage == "full" else ["S01", "S02"])
    prepared = bridge.prepare(order, stage, selected, draft=draft)
    workspace = Path(prepared["workspace"])
    count = len(selected)
    (workspace / "assets" / "style.css").write_text("body { color: #123; }", encoding="utf-8")
    for number in range(1, count + 1):
        (workspace / "slides" / f"slide{number:02d}.html").write_text(
            "<!doctype html><html><body>SYNTHETIC TEST ONLY</body></html>", encoding="utf-8")
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    for _ in selected:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(2))
        box.text_frame.paragraphs[0].text = "SYNTHETIC INTEGRITY TEST ONLY"
        box.text_frame.paragraphs[0].font.size = Pt(24)
    deck.save(workspace / "answer.pptx")
    checksum = core.file_hash(workspace / "answer.pptx")
    renders, reviews = {}, {}
    for number, slide in enumerate(bridge._slides(workspace), start=1):
        relative = slide.relative_to(workspace).as_posix()
        image = workspace / "qa" / "slides" / f"slide{number:02d}.jpg"
        image.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (1280, 720), "white").save(image)
        render = {"inputs": bridge._snapshot(workspace, [slide]),
                  "images": {image.relative_to(workspace).as_posix(): core.file_hash(image)}}
        renders[relative] = render
        reviews[relative] = {"ok": True, "verdict": "pass", "summary": "Synthetic integrity fixture only",
                             "issues": [], "review_mode": "host", "render": copy.deepcopy(render)}
    contact = workspace / "qa" / "deck" / "contact-01.jpg"
    contact.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (1280, 720), "gray").save(contact)
    deck_render = {"inputs": {"artifact_sha256": checksum},
                   "images": {contact.relative_to(workspace).as_posix(): core.file_hash(contact)}}
    core.write_json(workspace / "workspace-state.json", {
        "schema_version": bridge.STATE_SCHEMA,
        "renders": {"slides": renders, "deck": deck_render},
        "reviews": {"slides": reviews, "deck": {
            "ok": True, "verdict": "pass", "summary": "Synthetic integrity fixture only",
            "issues": [], "review_mode": "host", "render": copy.deepcopy(deck_render)}},
        "build": {"ok": True, "artifact": "answer.pptx", "artifact_sha256": checksum,
                  "inputs": bridge._snapshot(workspace, bridge._slides(workspace)),
                  "slides": count, "aspect_ratio": "16:9", "quality_warnings": []},
    })
    core.write_json(workspace / "final-report.json", {
        "schema_version": bridge.FINAL_SCHEMA, "ok": True, "complete": True, "delivery_ready": True,
        "checks": {name: True for name in bridge.FINAL_CHECKS}, "artifact": "answer.pptx",
        "artifact_sha256": checksum, "slides": count, "aspect_ratio": "16:9",
    })
    return prepared, workspace


def imported_record(order, workspace):
    result = bridge.validate_result(order, workspace)
    return {**result, "path": str(Path(result["pptx_path"]).relative_to(order)),
            "sha256": core.file_hash(Path(result["pptx_path"]))}


def test_manifest_is_pinned_and_modified_source_is_rejected(tmp_path, monkeypatch):
    assert bridge.vendor_signature() == bridge.SOURCE_MANIFEST_SIGNATURE
    vendor = tmp_path / "vendor"
    shutil.copytree(bridge._vendor_root(), vendor)
    monkeypatch.setattr(bridge, "_vendor_root", lambda: vendor)
    (vendor / "SKILL.md").write_text("modified", encoding="utf-8")
    with pytest.raises(ValueError, match="source changed"):
        bridge.vendor_signature()
    manifest = core.read_json(vendor / "UPSTREAM.json")
    manifest["files"]["SKILL.md"] = core.file_hash(vendor / "SKILL.md")
    core.write_json(vendor / "UPSTREAM.json", manifest)
    with pytest.raises(ValueError, match="bundled digest"):
        bridge.vendor_signature()


def test_runtime_discovery_does_not_execute_commands(monkeypatch):
    monkeypatch.setattr(subprocess, "run", lambda *a, **k: pytest.fail("unexpected subprocess"))
    monkeypatch.setattr(subprocess, "Popen", lambda *a, **k: pytest.fail("unexpected subprocess"))
    monkeypatch.setattr(bridge.sys, "platform", "win32")
    info = bridge.runtime_info()
    assert info["revision"] == bridge.REVISION
    assert info["skill_path"].endswith("SKILL.md")
    assert info["supported_platform"] is False
    assert info["ready"] is False
    assert info["automatic_install"] is False


def test_setup_check_and_windows_install_do_not_mutate_or_execute(monkeypatch, capsys):
    spec = importlib.util.spec_from_file_location("setup_pptagent", Path(__file__).parents[1] / "scripts" / "setup_pptagent.py")
    setup = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(setup)
    monkeypatch.setattr(subprocess, "run", lambda *a, **k: pytest.fail("unexpected subprocess"))
    monkeypatch.setattr(setup.venv.EnvBuilder, "create", lambda *a, **k: pytest.fail("unexpected install"))
    monkeypatch.setattr(bridge.sys, "platform", "win32")
    assert setup.main(["--check"]) == 0
    assert json.loads(capsys.readouterr().out)["status"] == "check_only"
    assert setup.main(["--install"]) == 2
    assert json.loads(capsys.readouterr().out)["status"] == "unsupported_platform"


def test_prepare_registers_fresh_requests_and_never_claims_generation_or_approval(order):
    before = core.model_digest(order)
    first = bridge.prepare(order, "sample", ["S01", "S03"])
    second = bridge.prepare(order, "sample", ["S01", "S03"])
    assert first["workspace"] != second["workspace"]
    assert first["status"] == "needs_authoring"
    assert first["customer_approval_created"] is False
    workspace = Path(first["workspace"])
    assert not (workspace / "answer.pptx").exists()
    assert not (workspace / "workspace-state.json").exists()
    assert not (workspace / "final-report.json").exists()
    assert core.read_json(workspace / "task.json")["slides"] == 2
    state = core.read_json(order / "state.json")
    assert len(state["design_requests"]) == 2
    assert state["approvals"] == {}
    assert core.model_digest(order) == before


def test_sample_handoff_limits_facts_and_sources_to_page_references(order):
    core.write_json(order / "facts.json", {"facts": [
        {"id": "F1", "text": "Included", "source_refs": ["SRC1#line1"]},
        {"id": "F2", "text": "Unrelated", "source_refs": ["SRC1#line2"]},
        {"id": "F3", "text": "Unrelated", "source_refs": ["SRC2"]},
    ]})
    core.write_json(order / "sources.json", {"files": [
        {"id": "SRC1", "path": "sources/one.md", "status": "ready"},
        {"id": "SRC2", "path": "sources/two.md", "status": "ready"},
    ]})
    deck = core.read_json(order / "deck.json")
    deck["slides"][0]["source_refs"] = ["F1"]
    deck["slides"][1]["source_refs"] = ["F3"]
    core.write_json(order / "deck.json", deck)
    handoff = bridge.prepare(order, "sample", ["S01"])
    content = core.read_json(Path(handoff["workspace"]) / "handoff-content.json")
    assert [fact["id"] for fact in content["facts"]["facts"]] == ["F1"]
    assert [source["id"] for source in content["source_index"]] == ["SRC1"]
    assert content["slides"][0]["source_refs"] == ["F1"]


def test_complete_synthetic_contract_can_be_verified_and_does_not_approve(order):
    prepared, workspace = fixture_result(order)
    record = imported_record(order, workspace)
    assert record["request_id"] == prepared["request_id"]
    assert record["slide_ids"] == ["S01", "S02"]
    assert record["external_evidence"]["review_basis"].endswith("not_independent_visual_certification")
    bridge.verify_evidence(order, record)
    assert core.read_json(order / "state.json")["approvals"] == {}


def test_final_report_alone_is_not_accepted(order):
    _, workspace = fixture_result(order)
    (workspace / "workspace-state.json").unlink()
    with pytest.raises(ValueError, match="missing"):
        bridge.validate_result(order, workspace)


@pytest.mark.parametrize("relative", [
    "slides/slide01.html", "assets/style.css", "qa/slides/slide01.jpg", "qa/deck/contact-01.jpg", "answer.pptx",
])
def test_changed_sources_renders_or_pptx_invalidate_current_and_imported_evidence(order, relative):
    _, workspace = fixture_result(order)
    record = imported_record(order, workspace)
    with (workspace / relative).open("ab") as stream:
        stream.write(b"changed")
    with pytest.raises(ValueError):
        bridge.validate_result(order, workspace)
    with pytest.raises(ValueError):
        bridge.verify_evidence(order, record)


def test_new_shared_asset_invalidates_saved_evidence(order):
    _, workspace = fixture_result(order)
    record = imported_record(order, workspace)
    (workspace / "assets" / "new.css").write_text("/* new dependency */", encoding="utf-8")
    with pytest.raises(ValueError, match="sources changed"):
        bridge.verify_evidence(order, record)


def test_forged_complete_report_cannot_hide_missing_reviews_or_stale_snapshots(order):
    _, workspace = fixture_result(order)
    state = core.read_json(workspace / "workspace-state.json")
    state["reviews"]["slides"].clear()
    core.write_json(workspace / "workspace-state.json", state)
    with pytest.raises(ValueError, match="visual review"):
        bridge.validate_result(order, workspace)


def test_hashed_nonimage_cannot_masquerade_as_a_review_render(order):
    _, workspace = fixture_result(order)
    image = workspace / "qa" / "deck" / "contact-01.jpg"
    image.write_text("This is text, not a rendered deck", encoding="utf-8")
    state = core.read_json(workspace / "workspace-state.json")
    new_hash = core.file_hash(image)
    state["renders"]["deck"]["images"]["qa/deck/contact-01.jpg"] = new_hash
    state["reviews"]["deck"]["render"]["images"]["qa/deck/contact-01.jpg"] = new_hash
    core.write_json(workspace / "workspace-state.json", state)
    with pytest.raises(ValueError, match="readable image"):
        bridge.validate_result(order, workspace)


@pytest.mark.parametrize("kind", ["failed_review", "escape_path", "wrong_count", "malformed_state"])
def test_rejects_invalid_proof_independent_of_report_flags(order, kind):
    _, workspace = fixture_result(order)
    state = core.read_json(workspace / "workspace-state.json")
    if kind == "failed_review":
        state["reviews"]["deck"]["verdict"] = "fail"
    elif kind == "escape_path":
        state["renders"]["deck"]["images"] = {"../outside.jpg": "0" * 64}
    elif kind == "wrong_count":
        state["build"]["slides"] = 99
    else:
        state["reviews"] = []
    core.write_json(workspace / "workspace-state.json", state)
    with pytest.raises(ValueError):
        bridge.validate_result(order, workspace)


def test_changed_order_or_old_round_result_cannot_be_imported(order):
    _, workspace = fixture_result(order)
    original = core.read_json(order / "style.json")
    core.write_json(order / "style.json", {**original, "accent": "000000"})
    with pytest.raises(ValueError, match="Order input changed"):
        bridge.validate_result(order, workspace)
    core.write_json(order / "style.json", original)
    record = imported_record(order, workspace)
    switch_design(order, "pptagent", "Second synthetic round")
    with pytest.raises(ValueError, match="previous design round"):
        bridge.validate_result(order, workspace)
    with pytest.raises(ValueError, match="another engine/design round"):
        bridge.verify_evidence(order, record)


def test_full_handoff_references_accepted_sample_in_separate_workspace(order):
    _, sample_workspace = fixture_result(order)
    sample = imported_record(order, sample_workspace)
    state = core.read_json(order / "state.json")
    state["artifacts"]["sample"] = sample
    core.write_json(order / "state.json", state)
    full = bridge.prepare(order, "full", ["S01", "S02", "S03"])
    assert full["workspace"] != str(sample_workspace)
    reference = core.read_json(Path(full["workspace"]) / "handoff-content.json")["approved_sample"]
    assert reference["workspace"] == sample["external_evidence"]["workspace"]
    assert Path(reference["pptx_path"]).is_file()
    assert reference["external_evidence"]["request_id"] == sample["request_id"]
    assert "reuse the approved sample" in full["instructions"]
    bridge.verify_evidence(order, sample)


def test_non_sample_copy_edit_does_not_destroy_saved_sample_evidence(order):
    _, workspace = fixture_result(order, ["S01"])
    sample = imported_record(order, workspace)
    deck = core.read_json(order / "deck.json")
    deck["slides"][2]["body"] = ["New full-deck copy only"]
    core.write_json(order / "deck.json", deck)
    bridge.verify_evidence(order, sample)
    with pytest.raises(ValueError, match="Order input changed"):
        bridge.validate_result(order, workspace)


def test_windows_separator_evidence_is_portable(order):
    _, workspace = fixture_result(order)
    state = core.read_json(workspace / "workspace-state.json")
    state["build"]["inputs"]["files"] = {key.replace("/", "\\"): value for key, value in state["build"]["inputs"]["files"].items()}
    for category in ("renders", "reviews"):
        state[category]["slides"] = {key.replace("/", "\\"): value for key, value in state[category]["slides"].items()}
        for record in state[category]["slides"].values():
            render = record if category == "renders" else record["render"]
            render["inputs"]["files"] = {key.replace("/", "\\"): value for key, value in render["inputs"]["files"].items()}
            render["images"] = {key.replace("/", "\\"): value for key, value in render["images"].items()}
    core.write_json(workspace / "workspace-state.json", state)
    bridge.validate_result(order, workspace)
